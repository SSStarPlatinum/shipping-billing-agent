"""Excel / Word 解析单测（现场生成文件，避免仓库里放二进制样例）。"""

from __future__ import annotations

from pathlib import Path

from shipping_billing.extract.document_extractor import DocumentExtractor
from shipping_billing.ingest import OcrClient, load_document


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "费用明细"
    rows = [
        ["单据类型", "Debit Note"],
        ["单据编号", "DN-202610-001"],
        ["单据日期", "2026-10-20"],
        ["币种", "USD"],
        ["签发方", "NEPTUNE SHIPPING LIMITED"],
        ["接收方", "CIMC WETRANS INTL LOG CO LTD"],
        [],
        ["费用名称", "数量", "单价", "金额", "币种", "付款方式"],
        ["Terminal Handling Charge", 2, 150.0, 300.0, "USD", "PP"],
        ["Doc Fee", 1, 50.0, 50.0, "USD", "PP"],
        ["合计", "", "", 350.0, "", ""],
    ]
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("DEBIT NOTE", level=1)
    document.add_paragraph("单据编号：DN-202611-002")
    document.add_paragraph("单据日期：2026-11-05")
    document.add_paragraph("币种：USD")
    document.add_paragraph("签发方：NEPTUNE SHIPPING LIMITED")
    document.add_paragraph("接收方：CIMC WETRANS INTL LOG CO LTD")
    document.add_paragraph("付款方式：CC")
    table = document.add_table(rows=1, cols=5)
    header = table.rows[0].cells
    for index, text in enumerate(["费用名称", "数量", "单价", "金额", "币种"]):
        header[index].text = text
    for name, qty, rate, amount in [
        ("Storage Fee", 3, 25.0, 75.0),
        ("Trucking Fee", 1, 320.0, 320.0),
    ]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = str(qty)
        cells[2].text = f"{rate:.2f}"
        cells[3].text = f"{amount:.2f}"
        cells[4].text = "USD"
    document.save(path)


def test_excel_fee_table(settings, tmp_path: Path) -> None:
    path = tmp_path / "fees.xlsx"
    _make_xlsx(path)
    raw = load_document(path, ocr_client=OcrClient(settings))
    assert raw.kind == "excel"
    assert raw.tables and raw.tables[0].name == "费用明细"

    doc = DocumentExtractor(settings).extract_one(raw)
    assert doc.document_type == "Debit Note"
    assert doc.reference_no == "DN-202610-001"
    assert doc.document_date == "2026-10-20"
    assert doc.currency == "USD"
    assert doc.payment_term == "PP"
    assert doc.issuer == "NEPTUNE SHIPPING LIMITED"
    assert doc.recipient == "CIMC WETRANS INTL LOG CO LTD"
    names = [c.charge_item for c in doc.charges]
    amounts = [c.amount for c in doc.charges]
    assert names == ["Terminal Handling Charge", "Documentation Fee"]
    assert amounts == [300.0, 50.0]
    assert doc.charges[0].quantity == 2
    assert doc.charges[0].unit_rate == 150.0


def test_word_fee_table(settings, tmp_path: Path) -> None:
    path = tmp_path / "fees.docx"
    _make_docx(path)
    raw = load_document(path, ocr_client=OcrClient(settings))
    assert raw.kind == "word"
    assert raw.tables

    doc = DocumentExtractor(settings).extract_one(raw)
    assert doc.document_type == "Debit Note"
    assert doc.reference_no == "DN-202611-002"
    assert doc.payment_term == "CC"
    names = [c.charge_item for c in doc.charges]
    amounts = [c.amount for c in doc.charges]
    assert names == ["Storage Charge", "Trucking Charge"]
    assert amounts == [75.0, 320.0]
