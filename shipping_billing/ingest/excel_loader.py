"""Excel（.xlsx/.xlsm/.xls）解析：读取全部工作表与单元格，保留二维表格。"""

from __future__ import annotations

from pathlib import Path

from ..io_utils import get_logger, normalize_text, sha256_bytes
from ..models import RawDocument, RawTable
from .ocr_client import OcrClient

logger = get_logger("excel")

EXCEL_EXTS = {".xlsx", ".xlsm", ".xltx", ".xls"}


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.10f}".rstrip("0").rstrip(".")
    return str(value).strip()


def _load_xlsx(path: Path) -> tuple[list[RawTable], str]:
    from openpyxl import load_workbook

    tables: list[RawTable] = []
    lines: list[str] = []
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_to_str(c) for c in row]
                while values and not values[-1]:
                    values.pop()
                if not values:
                    continue
                rows.append(values)
            if not rows:
                continue
            tables.append(RawTable(name=sheet.title, locator=f"sheet:{sheet.title}", rows=rows))
            lines.append(f"[工作表 {sheet.title}]")
            for row in rows:
                lines.append(" | ".join(row))
    finally:
        workbook.close()
    return tables, "\n".join(lines)


def _load_xls(path: Path) -> tuple[list[RawTable], str]:
    try:
        import xlrd  # 可选依赖
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("读取 .xls 需要安装 xlrd（pip install xlrd），或先另存为 .xlsx") from exc

    book = xlrd.open_workbook(filename=str(path))
    tables: list[RawTable] = []
    lines: list[str] = []
    for sheet in book.sheets():
        rows: list[list[str]] = []
        for r in range(sheet.nrows):
            values = [_cell_to_str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            while values and not values[-1]:
                values.pop()
            if values:
                rows.append(values)
        if not rows:
            continue
        tables.append(RawTable(name=sheet.name, locator=f"sheet:{sheet.name}", rows=rows))
        lines.append(f"[工作表 {sheet.name}]")
        for row in rows:
            lines.append(" | ".join(row))
    return tables, "\n".join(lines)


def parse_excel(path: Path, *, ocr_client: OcrClient, role: str = "primary") -> RawDocument:
    raw_bytes = path.read_bytes()
    sha = sha256_bytes(raw_bytes)
    suffix = path.suffix.lower()
    warnings: list[str] = []
    try:
        tables, text = _load_xls(path) if suffix == ".xls" else _load_xlsx(path)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Excel 解析失败 %s: %s", path, exc)
        tables, text = [], ""
        warnings.append(f"Excel 解析失败: {type(exc).__name__}: {exc}")

    text = normalize_text(text)
    if not text:
        warnings.append("Excel 未读取到有效内容")

    return RawDocument(
        doc_id=sha[:12],
        path=str(path),
        filename=path.name,
        kind="excel",
        sha256=sha,
        size=len(raw_bytes),
        text=text,
        pages=[text] if text else [],
        tables=tables,
        metadata={"sheet_count": len(tables), "sheets": [t.name for t in tables]},
        warnings=warnings,
    )


__all__ = ["parse_excel", "EXCEL_EXTS"]
