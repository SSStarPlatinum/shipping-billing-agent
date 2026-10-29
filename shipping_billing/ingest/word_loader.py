"""Word（.docx/.doc）解析：正文段落 + 表格。"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from ..io_utils import get_logger, normalize_text, sha256_bytes
from ..models import RawDocument, RawTable
from .ocr_client import OcrClient

logger = get_logger("word")

WORD_EXTS = {".docx", ".docm", ".doc"}


def _load_docx(path: Path) -> tuple[list[RawTable], str]:
    from docx import Document

    document = Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        if para.text and para.text.strip():
            lines.append(para.text.strip())

    tables: list[RawTable] = []
    for index, table in enumerate(document.tables, start=1):
        rows: list[list[str]] = []
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            while values and not values[-1]:
                values.pop()
            if values:
                rows.append(values)
        if rows:
            tables.append(RawTable(name=f"table-{index}", locator=f"table {index}", rows=rows))
            lines.append(f"[表格 {index}]")
            for row in rows:
                lines.append(" | ".join(row))
    return tables, "\n".join(lines)


def _convert_with_libreoffice(path: Path, work_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    out_dir = Path(tempfile.mkdtemp(dir=str(work_dir)))
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", str(out_dir), str(path)],
            check=True,
            capture_output=True,
            timeout=180,
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("LibreOffice 转换失败: %s", exc)
        return None
    converted = out_dir / (path.stem + ".docx")
    return converted if converted.exists() else None


def _naive_doc_text(path: Path) -> str:
    """旧版 .doc 兜底：提取可打印字符串（不保证版式与表格边界）。"""
    raw = path.read_bytes()
    chunks = re.findall(rb"[\x20-\x7e]{4,}", raw)
    text = " ".join(c.decode("latin-1", errors="ignore") for c in chunks)
    # 去掉 OLE 结构噪声
    text = re.sub(r"[A-Za-z]:\\\\?[^\s]{0,80}", " ", text)
    return normalize_text(text)


def parse_word(
    path: Path, *, ocr_client: OcrClient, role: str = "primary", work_dir: Path | None = None
) -> RawDocument:
    raw_bytes = path.read_bytes()
    sha = sha256_bytes(raw_bytes)
    warnings: list[str] = []
    tables: list[RawTable] = []
    text = ""

    suffix = path.suffix.lower()
    if suffix == ".doc":
        work = Path(work_dir or tempfile.gettempdir())
        converted = _convert_with_libreoffice(path, work)
        if converted is not None:
            try:
                tables, text = _load_docx(converted)
                warnings.append("旧版 .doc 已通过 LibreOffice 转换后解析")
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"转换后的 .docx 解析失败: {exc}")
        if not text:
            text = _naive_doc_text(path)
            if text:
                warnings.append(
                    "旧版 .doc 使用二进制兜底提取，文字可能存在噪声，建议转为 .docx"
                )
            else:
                warnings.append("无法解析旧版 .doc，请转换为 .docx")
    else:
        try:
            tables, text = _load_docx(path)
        except Exception as exc:  # noqa: BLE001
            logger.warning("Word 解析失败 %s: %s", path, exc)
            warnings.append(f"Word 解析失败: {type(exc).__name__}: {exc}")

    text = normalize_text(text)
    if not text:
        warnings.append("Word 未读取到有效内容")

    return RawDocument(
        doc_id=sha[:12],
        path=str(path),
        filename=path.name,
        kind="word",
        sha256=sha,
        size=len(raw_bytes),
        text=text,
        pages=[text] if text else [],
        tables=tables,
        metadata={"table_count": len(tables)},
        warnings=warnings,
    )


__all__ = ["parse_word", "WORD_EXTS"]
